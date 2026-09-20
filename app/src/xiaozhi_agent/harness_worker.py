from __future__ import annotations

import json
import logging
import os
import shutil
from pathlib import Path
from typing import Callable

from .config import AppConfig
from .models import WorkResult
from .paths import DSH_HOME, PRIVATE_PYTHON, RUNTIME_DIR, SKILLS_DIR

log = logging.getLogger(__name__)


class HarnessWorker:
    def __init__(self, config: AppConfig, api_key: str):
        self.config = config
        self.api_key = api_key
        workspace = config.workspace_path
        if not workspace:
            raise RuntimeError("Harness workspace 未配置")
        self.workspace = workspace
        self.workspace.mkdir(parents=True, exist_ok=True)

    def deploy_workspace_support(self) -> None:
        """Install only project instructions/skills into the selected workspace."""
        skills_target = self.workspace / ".dsh" / "skills"
        skills_target.mkdir(parents=True, exist_ok=True)
        if SKILLS_DIR.exists():
            for child in SKILLS_DIR.iterdir():
                dst = skills_target / child.name
                if child.is_dir():
                    shutil.copytree(child, dst, dirs_exist_ok=True)
                else:
                    shutil.copy2(child, dst)

        agents = self.workspace / "AGENTS.md"
        runtime_python = PRIVATE_PYTHON if PRIVATE_PYTHON.exists() else Path(os.sys.executable)
        text = f"""# XiaoZhi workspace instructions

You are the complex-task worker behind XiaoZhi Task Center.

## Non-negotiable rules
- Work only inside the task directory supplied in the user's prompt unless reading source input requires otherwise.
- Never send messages to WeChat contacts yourself. Delivery/approval belongs to Task Center.
- Never overwrite an approved or previous output version.
- Final deliverables must be written under the exact output directory specified by the task prompt.
- Re-open and verify every final Office file before declaring completion.
- For Excel/Word/PPT generation and validation, use deterministic Python code where practical.

## Fixed Python runtime
Use exactly:
`{runtime_python}`

Installed libraries include pandas, openpyxl, XlsxWriter, python-docx, python-pptx, Pillow, matplotlib and requests.

## Office quality
Use the skills under `.dsh/skills/office-docx`, `.dsh/skills/office-xlsx`, `.dsh/skills/office-pptx` when relevant.
Prefer professional Chinese office formatting, consistent fonts, safe glyphs, sensible widths/heights, no PPT overflow, and structural validation.
"""
        agents.write_text(text, encoding="utf-8")

    def _task_dirs(self, task_id: str, version: int) -> dict[str, Path]:
        root = self.workspace / ".xiaozhi_tasks" / task_id
        dirs = {
            "root": root,
            "input": root / "input",
            "work": root / "work",
            "output": root / "versions" / f"v{version}",
        }
        for p in dirs.values():
            p.mkdir(parents=True, exist_ok=True)
        return dirs

    @staticmethod
    def _copy_inputs(files: list[Path], input_dir: Path) -> list[Path]:
        copied: list[Path] = []
        for src in files:
            src = Path(src)
            if not src.exists() or not src.is_file():
                continue
            dst = input_dir / src.name
            if src.resolve() != dst.resolve():
                shutil.copy2(src, dst)
            copied.append(dst)
        return copied

    def _env(self, mode: str) -> dict[str, str]:
        old_path = os.environ.get("PATH", "")
        private = str(RUNTIME_DIR)
        scripts = str(RUNTIME_DIR / "Scripts")
        return {
            "DSH_PERMISSION_MODE": mode,
            "PATH": private + os.pathsep + scripts + os.pathsep + old_path,
            "PYTHONUTF8": "1",
        }

    def _run_once(
        self,
        *,
        task_id: str,
        prompt: str,
        session_id: str,
        mode: str,
        on_event: Callable[[str], None] | None,
    ):
        from deepseek_harness import DeepSeekHarness

        with DeepSeekHarness(
            provider="deepseek-official",
            model=self.config.model,
            max_tokens=49_152,
            cwd=str(self.workspace),
            runtime_cwd=str(self.workspace),
            dsh_home=str(DSH_HOME),
            profile="sdk",
            api_key=self.api_key,
            base_url=self.config.deepseek_base_url,
            env=self._env(mode),
            initialize_timeout_seconds=45,
            request_timeout_seconds=None,
        ) as harness:
            def notify(note) -> None:
                if on_event is None:
                    return
                try:
                    method = getattr(note, "method", "")
                    payload = getattr(note, "payload", {})
                    if method == "session.event":
                        event = payload.get("event") or {}
                        etype = event.get("type")
                        if etype:
                            on_event(str(etype))
                except Exception:
                    return

            return harness.run(prompt, session_id=session_id, on_notification=notify)

    @staticmethod
    def _looks_like_windows_sandbox_bug(exc: BaseException | str) -> bool:
        text = str(exc).lower()
        signals = (
            "--profile <name> is required",
            "sandbox_unavailable",
            "windows-acl",
            "pwsh-sandbox",
        )
        return any(x in text for x in signals)

    def run_task(
        self,
        *,
        task_id: str,
        text: str,
        input_files: list[Path],
        next_version: int,
        session_id: str,
        revision_of: list[Path] | None = None,
        on_event: Callable[[str], None] | None = None,
    ) -> WorkResult:
        self.deploy_workspace_support()
        dirs = self._task_dirs(task_id, next_version)
        copied = self._copy_inputs(input_files, dirs["input"])
        previous = [Path(p) for p in (revision_of or []) if Path(p).exists()]

        prompt = f"""You are executing XiaoZhi task {task_id}.

USER REQUEST:
{text}

TASK DIRECTORY: {dirs['root']}
INPUT DIRECTORY: {dirs['input']}
WORK DIRECTORY: {dirs['work']}
FINAL OUTPUT DIRECTORY FOR THIS VERSION: {dirs['output']}
VERSION: V{next_version}
FIXED PYTHON: {PRIVATE_PYTHON}

INPUT FILES:
{json.dumps([str(p) for p in copied], ensure_ascii=False, indent=2)}

PREVIOUS VERSION FILES (if this is a revision; read them but DO NOT overwrite them):
{json.dumps([str(p) for p in previous], ensure_ascii=False, indent=2)}

Requirements:
1. Plan and perform the task autonomously. You may write/run code and repair failures.
2. Use web/research tools when the user request requires current external research.
3. Put only final user deliverables in FINAL OUTPUT DIRECTORY. Scratch code belongs in WORK DIRECTORY.
4. If this is a revision, create a new V{next_version}; never modify the previous-version files in place.
5. Re-open/validate final Excel/Word/PPT outputs with the installed Python libraries and available Office skills/checker.
6. Finish with a concise Chinese summary listing the final output paths. Do NOT send any file or message to a WeChat contact; Task Center handles delivery.
"""

        modes = ["workspace-write"] if self.config.harness_safe_mode else ["danger-full-access"]
        if self.config.harness_safe_mode and self.config.allow_harness_full_access_fallback:
            modes.append("danger-full-access")

        last_exc: Exception | None = None
        used_mode = modes[0]
        for idx, mode in enumerate(modes):
            used_mode = mode
            try:
                result = self._run_once(
                    task_id=task_id,
                    prompt=prompt,
                    session_id=session_id,
                    mode=mode,
                    on_event=on_event,
                )
                finish = getattr(result, "finish_reason", None)
                final_response = getattr(result, "final_response", "") or ""
                if finish == "error":
                    raise RuntimeError(f"Harness finish_reason=error: {final_response[:1200]}")
                outputs = self._collect_outputs(dirs["output"])
                verified, problems = self._verify_outputs(outputs)
                if not verified and outputs:
                    return WorkResult(
                        False,
                        text=final_response,
                        files=outputs,
                        error="; ".join(problems),
                        metadata={"harness_mode": mode, "finish_reason": finish},
                    )
                return WorkResult(
                    True,
                    text=final_response,
                    files=outputs,
                    metadata={"harness_mode": mode, "finish_reason": finish},
                )
            except Exception as exc:
                last_exc = exc
                log.exception("Harness run failed mode=%s task=%s", mode, task_id)
                if idx + 1 < len(modes) and self._looks_like_windows_sandbox_bug(exc):
                    if on_event:
                        on_event("windows_sandbox_bug_fallback_to_full_access")
                    continue
                break

        return WorkResult(
            False,
            error=str(last_exc or "Harness 运行失败"),
            metadata={"harness_mode": used_mode},
        )

    @staticmethod
    def _collect_outputs(output_dir: Path) -> list[Path]:
        return sorted([p for p in output_dir.rglob("*") if p.is_file() and p.stat().st_size > 0])

    @staticmethod
    def _verify_outputs(files: list[Path]) -> tuple[bool, list[str]]:
        problems: list[str] = []
        for p in files:
            try:
                suffix = p.suffix.lower()
                if suffix == ".xlsx":
                    from openpyxl import load_workbook
                    wb = load_workbook(p, read_only=True, data_only=False)
                    wb.close()
                elif suffix == ".docx":
                    from docx import Document
                    Document(p)
                elif suffix == ".pptx":
                    from pptx import Presentation
                    Presentation(p)
                elif suffix in {".html", ".htm", ".md", ".txt", ".csv"}:
                    p.read_bytes()
            except Exception as exc:
                problems.append(f"{p.name}: {exc}")
        return not problems, problems
